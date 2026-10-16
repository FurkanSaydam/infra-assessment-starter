"""Generate summary JSON, DOCX and PDF reports."""
import json
from pathlib import Path
from xml.etree import ElementTree as ET
from datetime import datetime, timezone

def write_reports(result_dir, results, targets, ports):
    result_dir = Path(result_dir)
    open_rows = []
    for result in results:
        for host in result.get("hosts", []):
            for port in host.get("ports", []):
                if port.get("state") == "open":
                    open_rows.append({"target": result.get("target"), **port})
    summary = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "target_count": len(targets),
        "completed_count": sum(r.get("status") == "completed" for r in results),
        "incomplete_count": sum(r.get("status") != "completed" for r in results),
        "open_port_count": len(open_rows), "targets": targets, "ports_requested": ports or "top 1000 TCP",
        "open_ports": open_rows, "results": results
    }
    (result_dir / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    # Combined XML: one well-formed nmaprun document containing host nodes from each scan.
    combined = ET.Element("nmaprun", {"scanner":"nmap", "args":"MFS Scanner combined results"})
    for result in results:
        xml_path = result.get("_xml_path")
        if xml_path and Path(xml_path).is_file():
            try:
                source = ET.parse(xml_path).getroot()
                for host in source.findall("host"):
                    combined.append(host)
            except (ET.ParseError, OSError):
                pass
    ET.ElementTree(combined).write(result_dir / "logs" / "nmap.xml", encoding="utf-8", xml_declaration=True)
    lines = []
    for result in results:
        lines.append(f'TARGET: {result.get("target")} STATUS: {result.get("status")}')
        for host in result.get("hosts", []):
            for p in host.get("ports", []):
                if p.get("state") == "open":
                    lines.append(f'  {p.get("port")}/{p.get("protocol")} {p.get("service") or "unknown"} {p.get("product") or ""} {p.get("version") or ""}'.rstrip())
    (result_dir / "logs" / "nmap.txt").write_text("\\n".join(lines) + "\\n", encoding="utf-8")
    (result_dir / "logs" / "web.txt").write_text("\\n".join(
        f'{r.get("target")}: {item}' for r in results for h in r.get("hosts", []) for item in h.get("web_review", [])) + "\\n", encoding="utf-8")
    (result_dir / "logs" / "smb.txt").write_text("\\n".join(
        f'{r.get("target")}: {item}' for r in results for h in r.get("hosts", []) for item in h.get("smb_review", [])) + "\\n", encoding="utf-8")
    (result_dir / "logs" / "rpc.txt").write_text("RPC-specific checks are not implemented in this version.\\n", encoding="utf-8")
    # DOCX/PDF dependencies are listed in requirements; keep JSON/log outputs even if absent.
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("MFS Scanner — Infrastructure Assessment", 0)
        doc.add_paragraph(f"Targets: {len(targets)} | Completed: {summary['completed_count']} | Incomplete: {summary['incomplete_count']}")
        doc.add_paragraph(f"Requested ports: {ports or 'Top 1000 TCP'}")
        doc.add_heading("Open ports", level=1)
        if open_rows:
            table = doc.add_table(rows=1, cols=5); table.style = "Light Shading Accent 1"
            for cell, value in zip(table.rows[0].cells, ["Target","Port","Protocol","Service","Version"]): cell.text=value
            for row in open_rows:
                cells = table.add_row().cells
                vals = [row.get("target"), row.get("port"), row.get("protocol"), row.get("service"), " ".join(str(row.get(k) or "") for k in ("product","version")).strip()]
                for cell, value in zip(cells, vals): cell.text=str(value or "")
        else: doc.add_paragraph("No open ports were parsed from the scan results.")
        doc.add_heading("Limitations", level=1)
        for text in ["Open ports and detected services do not by themselves prove a vulnerability.",
                     "Service version alone does not confirm CVE impact.",
                     "Only the configured targets and ports were assessed."]:
            doc.add_paragraph(text, style="List Bullet")
        doc.save(result_dir / "report.docx")
    except ImportError:
        pass
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
        pdf = SimpleDocTemplate(str(result_dir / "report.pdf"), pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        story = [Paragraph("MFS Scanner — Infrastructure Assessment", styles["Title"]), Spacer(1, 12),
                 Paragraph(f"Targets: {len(targets)} | Completed: {summary['completed_count']} | Incomplete: {summary['incomplete_count']}", styles["Normal"]),
                 Spacer(1, 12)]
        data = [["Target","Port","Protocol","Service","Version"]]
        data += [[str(x.get("target","")), str(x.get("port","")), str(x.get("protocol","")),
                  str(x.get("service","")), " ".join(str(x.get(k) or "") for k in ("product","version")).strip()] for x in open_rows]
        if len(data) == 1: data.append(["—","—","—","No open ports parsed","—"])
        table = Table(data, repeatRows=1, colWidths=[110,50,60,130,260])
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#12343b")),
                                   ("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.4,colors.grey),
                                   ("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),8)]))
        story += [table, Spacer(1, 16), Paragraph("Limitations: open ports do not prove vulnerabilities; service version alone does not confirm CVE impact.", styles["Normal"])]
        pdf.build(story)
    except ImportError:
        pass
